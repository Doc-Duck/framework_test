import sqlite3
import pandas as pd
from transliterate import translit
import openpyxl
import re
import io
import string


from Httpapp import db, apps, manager


def new_table(file_name, sheetname):
    # ----------------Добавление таблицы в бд----------------
    old_name = sheetname
    ss = openpyxl.load_workbook(file_name)
    ss_sheet = ss.get_sheet_by_name(sheetname)
    ss_sheet.title = re.sub(r'[^A-z0-9]', '', translit(sheetname, language_code='ru', reversed=True))
    sheetname = ss_sheet.title
    ss.save(file_name)

    con = sqlite3.connect('Httpapp/items.db')
    wb = pd.read_excel(file_name, sheetname)
    for column in wb:
        trans = translit(column, language_code='ru', reversed=True)
        trans = re.sub(r'[^A-z0-9]', '', trans)
        wb.rename(columns={f'{column}': f'{trans}'}, inplace=True)
    wb.to_sql(sheetname, con)
    con.commit()
    con.close()

    ss = openpyxl.load_workbook(file_name)
    ss_sheet = ss.get_sheet_by_name(sheetname)
    ss_sheet.title = old_name
    sheetname = old_name
    ss.save(file_name)

    # ----------------Чтение exel файла----------------

    result = pd.read_excel(file_name, sheetname).columns.values
    counter = 0
    for it in result:
        result[counter] = translit(result[counter], language_code='ru', reversed=True)
        result[counter] = re.sub(r'[^A-z0-9]', '', result[counter])
        counter += 1
    xl = pd.ExcelFile(file_name).sheet_names
    index = xl.index(sheetname)
    print(index)
    xl = pd.ExcelFile(file_name).sheet_names[index]
    xl = translit(xl, language_code='ru', reversed=True)

    # ----------------Работа с файлом классов бд----------------

    trash = f'class {xl}(db.Model):\n\tindex = db.Column(primary_key=True)\n'
    i = 0
    with open('Httpapp/models.py', 'r') as f:
        old_data = f.read()
    while i < len(result):
        trash = trash + '\t' + result[i].replace("'", '') + ' = db.Column()\n'
        i += 1
    with open('Httpapp/models.py', 'w') as f:
        f.write(old_data + trash + '\n')

    # ----------------Создание html файла таблицы----------------

    name = 'Httpapp/templates/' + xl + '.html'
    td = ''
    for column in result:
        column = column.replace(' ', '')
        column = column.replace("'", '')
        td = td + f'<td style="text-align: center">{{{{ el.{column} }}}}</td>' + '\n' + '\t\t\t'
        i += 1
    print(td)
    with open(name, 'w') as f:
        f.write(f'{{% extends "base.html" %}}\n\
        \n\
        {{% block body %}}\n\
        \n\
        <table class="table table-striped container">\n\
          <thead>\n\t\
          </thead>\n\t\
          <tbody>\n\t\
          {{% for el in {xl} %}}\n\
            <tr>\n\
                {td}\n\
            </tr>\t\t\n\
          {{% endfor %}}\n\t\
          </tbody>\n\t\
        </table>\n\
        {{% endblock %}}')

    # ----------------Создание нового маршрута----------------

    my_file = open('Httpapp/roots.py', 'a+')
    text_for_file = f'\n\n\n@apps.route("/{xl}")\n@login_required\ndef {xl.lower()}():\n\t{xl.lower()} = {xl}.query.all()\n\treturn render_template("{xl}.html", {xl}={xl.lower()})'
    my_file.write(text_for_file)
    my_file.close()

    with open('Httpapp/roots.py', 'r') as f:
        old_data = f.read()
    new_data = old_data.replace('#list_add', f',{xl} #list_add')
    with open('Httpapp/roots.py', 'w') as f:
        f.write(new_data)

    # ----------------Добавление нового элемента списка в тэмплейт----------------

    with open('Httpapp/templates/base.html', 'r') as f:
        old_data = f.read()
    new_data = old_data.replace('<!--- add link --->', f'<li><a href="{ xl }" class="dropdown-item">{ xl.lower() }</a></li>\n\t\t\t\t\t\t<!--- add link --->')
    with open('Httpapp/templates/base.html', 'w') as f:
        f.write(new_data)


def add_table(table_name, columns):
    table_cols = ''
    for col in columns:
        table_cols += f'<td style="text-align: center">{{{{ el.{col} }}}}</td>\n\t\t\t\t'

    with io.open(f'Httpapp/templates/l_template.html', 'r') as a:
        old_data = a.read()
    old_data = old_data.replace('<!--- table_cols --->', table_cols)
    old_data = old_data.replace('<!--- table_name --->', table_name)
    old_data = old_data.replace('<!--- edit_table --->', f'{table_name}_edit')
    for col in columns:
        old_data = old_data.replace('<!--- add_inputs --->', f'<input class="form-control mt-3 me-3" name="{col}" placeholder="{col}" type="text">\n\t\t\t<!--- add_inputs --->')
    f = open(f'Httpapp/templates/{table_name}.html', 'a+')
    f.write(old_data)
    f.close()

    trash = f'class {string.capwords(table_name)}(db.Model):\n\tindex = db.Column(db.Integer,primary_key=True)\n'
    with open('Httpapp/models.py', 'r') as f:
        old_data = f.read()
    for col in columns:
        trash = trash + '\t' + col + ' = db.Column(db.String)\n'
    with open('Httpapp/models.py', 'w') as f:
        f.write(old_data + trash + '\n')
    db.create_all()

    reqs = ''
    reqs_com = ''
    for col in columns:
        reqs += f"""{col} = request.form['{col}']
        """
        reqs_com += f"{col}={col},"
    reqs_com = reqs_com.rstrip(reqs_com[-1])

    my_file = open('Httpapp/roots.py', 'a+')
    text_for_file = f'''@apps.route('/{table_name}', methods=['POST', 'GET'])
@login_required
def {table_name}():
    if request.method == 'POST':
        type = request.form['type']
        column_x = request.form['column-x']
        column_y = request.form['column-y']
        file_name = request.form['file-name']
        new_chart(file_name, type, column_x, column_y)
    {table_name} = {string.capwords(table_name)}.query.order_by({string.capwords(table_name)}.index).all()
    return render_template("{table_name}.html", {table_name}={table_name})
    

@apps.route('/adprod_{table_name}', methods=['POST'])
@login_required
def addprod_{table_name}():
    if request.method == 'POST':
        {reqs}
        items = {string.capwords(table_name)}({reqs_com})
        db.session.add(items)
        db.session.commit()
        return redirect('{table_name}')'''
    my_file.write(text_for_file)
    my_file.close()

    with open('Httpapp/roots.py', 'r') as f:
        old_data = f.read()
    new_data = old_data.replace('#list_add', f',{string.capwords(table_name)} #list_add')
    with open('Httpapp/roots.py', 'w') as f:
        f.write(new_data)








